"""Portable, offline drivers for three fixed public-data software exercises.

Only local authored orchestration is executed. Third-party scripts in archives
are never imported. Scientific payload stays in the supplied data/run folders.
"""
from __future__ import annotations

import argparse
from contextlib import redirect_stderr, redirect_stdout
import csv
from datetime import datetime, timezone
import hashlib
import io
import json
import math
import os
import re
from pathlib import Path
import subprocess
import sys
import time
import zipfile


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def save(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('x', encoding='utf-8') as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2, allow_nan=False)


def relative_path(root, name):
    name = Path(name)
    if name.is_absolute() or '..' in name.parts or not name.parts:
        raise ValueError('manifest file must be a safe relative path')
    path = (root / name).resolve()
    if not path.is_relative_to(root.resolve()):
        raise ValueError('manifest relative path resolves outside source root')
    return path


def prepare_run(manifest, source_root, run_dir):
    """Verify every source/member before creating a fresh output directory."""
    root, output = Path(source_root).resolve(), Path(run_dir).resolve()
    if output.exists():
        raise FileExistsError(f'Output already exists; choose a new run directory: {output}')
    sources, payloads = {}, {}
    for source in manifest['sources']:
        path = relative_path(root, source['path'])
        if digest(path) != source['sha256']:
            raise ValueError(f'SHA256 mismatch: {path}')
        sources[source['path']] = path
    identity=manifest.get('identity_metadata')
    identity_checks=[]
    if identity is not None:
        with zipfile.ZipFile(sources[identity['source']]) as archive:
            if archive.namelist().count(identity['member'])!=1:
                raise ValueError('identity metadata member must be unique')
            content=archive.read(identity['member'])
        if hashlib.sha256(content).hexdigest()!=identity['sha256']:
            raise ValueError('identity metadata SHA256 mismatch')
        original=list(csv.DictReader(io.StringIO(content.decode('utf-8-sig')),delimiter=';'))
        for record in manifest['records']:
            matches=[row for row in original if row.get('Non-redundant ID (sample ID)')==record['id']]
            if len(matches)!=1 or any(matches[0].get(k)!=v for k,v in record['context'].items()):
                raise ValueError('identity metadata mapping differs from original specimen fields')
            trial=re.fullmatch(r'execution_essai(\d+)',matches[0]['Filename'],re.IGNORECASE)
            if trial is None or int(trial[1])!=record['trial']:
                raise ValueError('identity metadata mapping differs from original trial filename')
            native_trial=re.search(r"d'essai (\d+) ",record.get('member',''))
            if native_trial is not None and int(native_trial[1])!=record['trial']:
                raise ValueError('identity metadata mapping differs from native member trial')
            identity_checks.append(dict(sample_id=record['id'],trial=record['trial'],
                metadata_source_row=original.index(matches[0])+2,fields_exact=True))
    if not manifest['records'] or len({r['id'] for r in manifest['records']}) != len(manifest['records']):
        raise ValueError('manifest needs distinct record identities')
    for record in manifest['records']:
        path = sources[record['source']]
        if record.get('member'):
            with zipfile.ZipFile(path) as archive:
                if archive.namelist().count(record['member']) != 1:
                    raise ValueError('archive must contain exactly one named source member')
                data = archive.read(record['member'])
        else:
            data = path.read_bytes()
        if hashlib.sha256(data).hexdigest() != record['sha256']:
            raise ValueError(f"SHA256 mismatch for record {record['id']}")
        relative_path(output / 'raw', record['output_name'])
        payloads[record['id']] = data
    if len({r['output_name'] for r in manifest['records']}) != len(manifest['records']):
        raise ValueError('record output filenames must be distinct')
    output.mkdir(parents=True, exist_ok=False)
    (output / 'raw').mkdir()
    paths = {}
    for record in manifest['records']:
        path = relative_path(output / 'raw', record['output_name'])
        path.parent.mkdir(parents=True, exist_ok=True)
        with path.open('xb') as stream:
            stream.write(payloads[record['id']])
        paths[record['id']] = path
    before = {str(path): digest(path) for path in sources.values()}
    save(output / 'source-hashes-before.json', before)
    save(output / 'frozen-manifest.json', manifest)
    if identity_checks:
        save(output/'specimen-identity-checks.json',dict(metadata=identity,checks=identity_checks))
    save(output / 'execution.json', {'started_at': datetime.now(timezone.utc).isoformat(),
        'case': manifest['case'], 'python': sys.executable, 'driver_sha256': digest(__file__),
        'argv': sys.argv, 'source_root': str(root), 'run_dir': str(output),
        'offline': True, 'third_party_code_executed': False})
    return dict(output=output, manifest=manifest, paths=paths, sources=sources,
                before=before, commands=[], readbacks=[])


def read_numeric(path, recipe):
    """Parse an explicit native layout without dropping rows or filling blanks."""
    if recipe['format'] == 'xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            sheet = book[recipe['sheet']]
            rows = list(sheet.iter_rows(values_only=True))
        finally:
            book.close()
    else:
        with Path(path).open(encoding=recipe['encoding'], newline='') as stream:
            rows = list(csv.reader(stream, delimiter=recipe['delimiter']))
    for check in recipe.get('header_checks', []):
        try:
            actual = rows[check['row']-1][check['column']]
        except (IndexError, TypeError) as exc:
            raise ValueError('missing required native header') from exc
        if str(actual).strip() != check['value']:
            raise ValueError(f"native header mismatch at row {check['row']}")
    first = recipe['start_row']
    data = rows[first-1:]
    if not data:
        raise ValueError('empty numeric block')
    width = len(recipe['fields'])
    if len(recipe['units']) != width or len(set(recipe['fields'])) != width:
        raise ValueError('field names and units must form an explicit unique layout')
    values = []
    for row in data:
        # Native CSV blocks must have exact width; XLSX may carry empty trailing
        # formatting columns, but populated cells beyond the declared block fail.
        if recipe['format'] == 'xlsx' and len(row) > width and all(x is None for x in row[width:]):
            row = row[:width]
        if len(row) != width or any(x is None or x == '' or isinstance(x, bool) for x in row):
            raise ValueError('empty cell or unexpected numeric block width')
        try:
            converted = [float(x) for x in row]
        except (TypeError, ValueError) as exc:
            raise ValueError('non-numeric native cell; formulas are not evaluated') from exc
        if not all(math.isfinite(x) for x in converted):
            raise ValueError('nonfinite native numeric cell')
        values.append(converted)
    return values, list(range(first, first+len(values)))


def invoke(context, name, command, config, destination):
    from experiment_to_cpfe.cli import main
    output = context['output']
    config_path = output / 'configs' / f'{name}.json'
    save(config_path, config)
    args = [command, '--config', str(config_path), '--run-dir', str(destination)]
    stdout, stderr = io.StringIO(), io.StringIO()
    started = time.perf_counter()
    print(f'开始：{name}', file=sys.stderr, flush=True)
    with redirect_stdout(stdout), redirect_stderr(stderr):
        code = main(args)
    receipt = dict(argv=['pipeline', *args], exit_code=code,
        seconds=time.perf_counter()-started, stdout=stdout.getvalue(), stderr=stderr.getvalue())
    context['commands'].append(receipt)
    save(output/'commands'/f"{len(context['commands']):02d}-{name}.json", receipt)
    if code:
        print(f'失败：{name}；查看 commands/ 中的诊断', file=sys.stderr, flush=True)
        raise RuntimeError(f'{name} rejected; see its saved command diagnostic')
    print(f"完成：{name}（{receipt['seconds']:.2f} 秒）", file=sys.stderr, flush=True)
    return destination


def normalize_record(context, record, selected=None, suffix=None):
    import h5py
    import numpy as np
    path, recipe = context['paths'][record['id']], record['reader']
    values, physical_rows = read_numeric(path, recipe)
    fields = recipe['fields'] if selected is None else selected
    indices = [recipe['fields'].index(name) for name in fields]
    units = {name:recipe['units'][i] for name,i in zip(fields,indices)}
    block = dict(data_start_row=physical_rows[0], data_end_row=physical_rows[-1]+1,
                 numeric_fields=fields, header_checks=recipe.get('header_checks', []))
    if recipe.get('sheet'):
        block['sheet'] = recipe['sheet']
    source = dict(path=str(path), table_name='measured_observations', source_kind='measured',
        modality='time_series', format=recipe['format'], delimiter=recipe.get('delimiter'),
        encoding=recipe.get('encoding', 'utf-8'), column_map={name:str(i) for name,i in zip(fields,indices)},
        units=units, coordinate_frame='instrument-record', axis_order=['row'],
        native_layout=record.get('layout_description', 'Explicit frozen instrument block; no value filtering'),
        license=context['manifest']['license'], block=block)
    config = dict(version=1, purpose=context['manifest'].get('purpose', 'Synthetic data-format exercise'),
        sample=dict(sample_id=record['id'], experiment_id=context['manifest']['doi'],
            microstructure_id='not-independently-characterized', load_path_id='source-recorded-instrument-program',
            schema_version='0.1', coordinate=dict(name='instrument-record',axes=['row'],units='1'),
            unit_system=dict(zip(recipe['fields'], recipe['units'])), tensor_order=['scalar'],
            orientation=dict(representation='not_applicable',reason='Scalar instrument-channel example')),
        sources=[source], solver_inputs=dict(source_identity=record.get('identity_evidence', 'synthetic record'),
            original_source={'file':record['source'], 'member':record.get('member'), 'sha256':record['sha256']},
            specimen_context=record.get('context', {}),
            limitations=context['manifest'].get('limitations', []),
            **({'dataset_split':record['role']} if record['role'] in {'train','validation','test'} else {})))
    label = record['id'] + (f'-{suffix}' if suffix else '')
    destination = context['output']/'normalized'/label
    invoke(context, 'normalize-'+label, 'normalize-sample', config, destination)
    with h5py.File(destination/'sample.h5','r') as handle:
        columns = handle['measured/measured_observations/columns']
        restored = np.column_stack([columns[name][:] for name in fields])
        np.testing.assert_array_equal(restored, np.asarray(values)[:,indices])
        np.testing.assert_array_equal(columns['source_row'][:],physical_rows)
        if recipe.get('sheet'):
            raw_records=json.loads(handle['measured/measured_observations/records_json'][()])
            if not all(row['source_sheet']==recipe['sheet'] for row in raw_records):
                raise ValueError('source sheet identity was not preserved')
    entry=dict(id=record['id'], role=record['role'], hdf5=str(destination/'sample.h5'),
        rows=len(values), selected_fields=fields, source_sha256=record['sha256'],
        numeric_readback_exact=True, physical_row_readback_exact=True)
    context['readbacks'].append(entry)
    return entry


def table_selector(column, unit, factor=None):
    result=dict(kind='table',table='measured_observations',column=column,
                source_unit=unit,id_columns=['source_row'])
    if factor is not None:
        result['conversion']=dict(factor=factor,offset=0.0,reason='Explicit percentage to dimensionless strain')
    return result


def select_candidate(candidates):
    if not candidates or any(not math.isfinite(c['validation_rmse']) for c in candidates):
        raise ValueError('candidate validation metrics must be finite')
    return min(candidates, key=lambda c:(c['validation_rmse'],c['id']!='linear'))


def grouped_build(context, entries, features, targets, selectors, *, sample_rows=False, scalar=False):
    import numpy as np
    layouts, inputs = {}, []
    for entry in entries:
        layout=dict(columns=selectors,alignment_evidence='Physical source_row within the named original file')
        if sample_rows:
            layout['rows']=dict(indices=np.linspace(0,entry['rows']-1,min(256,entry['rows']),dtype=int).tolist())
        layouts[entry['id']]=layout
        inputs.append(dict(path=entry['hdf5'],sample_id=entry['id'],layout=entry['id'],split=entry['role']))
    target_config=dict(version=1,target=targets[0]) if scalar else dict(version=2,targets=targets)
    config=dict(**target_config,features=features,group_by='sample_id',
        grouping_evidence=context['manifest'].get('grouping_evidence','File identities only; not batch independence'),
        layouts=layouts,inputs=inputs)
    return invoke(context,'build','build-training-dataset',config,context['output']/'dataset')


def run_dopamics(context):
    import numpy as np
    import torch
    torch.set_num_threads(1)
    records=context['manifest']['records']
    entries=[normalize_record(context,r) for r in records if r['role']!='test']
    test=next(r for r in records if r['role']=='test')
    features=[dict(name='laser_strain',unit='1')]
    targets=[dict(name='stress_reported',unit='MPa')]
    selectors=dict(laser_strain=table_selector('laser_strain_percent','%',.01),
                   stress_reported=table_selector('stress_reported','MPa'))
    dataset=grouped_build(context,entries,features,targets,selectors,sample_rows=True,scalar=True)
    budget=context['manifest']['training']
    candidates=[]
    for candidate in budget['candidates']:
        config=json.loads((dataset/'training-config.json').read_text(encoding='utf-8'))
        config.update(dataset=str(dataset/'dataset.npz'),architecture=candidate['architecture'],
            hidden=candidate['hidden'],seed=budget['seed'],epochs=budget['epochs'],
            patience=budget['patience'],learning_rate=budget['learning_rate'],evaluation_mode='external_test')
        out=invoke(context,'train-'+candidate['id'],'train-surrogate',config,context['output']/'models'/candidate['id'])
        receipt=json.loads((out/'training.json').read_text(encoding='utf-8'))
        candidates.append(dict(id=candidate['id'],path=str(out),model_sha256=digest(out/'model.pt'),
            validation_rmse=receipt['metrics']['validation']['rmse'],
            epochs_completed=receipt['epochs_completed'],best_epoch=receipt['best_epoch']))
    selected=select_candidate(candidates)
    save(context['output']/'model-selection.json',dict(candidates=candidates,selected=selected,
        selected_at=datetime.now(timezone.utc).isoformat(),rule='validation RMSE only; linear wins exact tie'))
    unlabeled=normalize_record(context,test,['laser_strain_percent'],'inputs-only')
    grouping=dict(group_by='sample_id',grouping_evidence=context['manifest']['grouping_evidence'])
    inference=dict(version=1,checkpoint=str(Path(selected['path'])/'model.pt'),
        checkpoint_sha256=selected['model_sha256'],features=features,**grouping,
        layouts={'curve':dict(columns={'laser_strain':selectors['laser_strain']},alignment_evidence='Physical source row')},
        inputs=[dict(path=unlabeled['hdf5'],sample_id=test['id'],layout='curve')])
    predicted=invoke(context,'infer','infer-surrogate',inference,context['output']/'inference')
    truth=normalize_record(context,test,['stress_reported'],'truth-only')
    evaluation=dict(version=1,predictions=str(predicted/'predictions.npz'),targets=targets,split='test',**grouping,
        layouts={'curve':dict(columns={'stress_reported':selectors['stress_reported']},alignment_evidence='Physical source row')},
        inputs=[dict(path=truth['hdf5'],sample_id=test['id'],layout='curve')])
    evaluated=invoke(context,'evaluate','evaluate-surrogate',evaluation,context['output']/'evaluation')
    from experiment_to_cpfe.learning.inference import read_predictions
    arrays,_,_=read_predictions(predicted/'predictions.npz')
    raw,physical_rows=read_numeric(context['paths'][test['id']],test['reader'])
    expected=np.asarray(raw)
    np.testing.assert_array_equal(arrays['features'][:,0],expected[:,7]*.01)
    if [json.loads(x)[0] for x in arrays['row_ids']] != physical_rows:
        raise ValueError('prediction row identity mismatch')
    with np.load(evaluated/'evaluation.npz',allow_pickle=False) as packed:
        np.testing.assert_array_equal(packed['targets'][:,0],expected[:,0])
        np.testing.assert_array_equal(packed['row_ids'],arrays['row_ids'])
    with np.load(dataset/'dataset.npz',allow_pickle=False) as packed:
        mean=float(np.mean(packed['targets'][packed['splits']=='train']))
    # A fresh Python process reads the saved model; never call any bundled third-party script.
    import experiment_to_cpfe
    backend_root=Path(experiment_to_cpfe.__file__).resolve().parent.parent
    program='import sys,json,numpy as np,experiment_to_cpfe; from pathlib import Path; from experiment_to_cpfe.learning.surrogate import predict_mlp; from experiment_to_cpfe.learning.inference import read_predictions; a,_,_=read_predictions(Path(sys.argv[1])); y=predict_mlp(Path(sys.argv[2]),a["features"]).reshape(-1,1); d=float(np.max(np.abs(y-a["prediction"]))); assert d==0; print(json.dumps({"max_abs_difference":d,"backend_root":str(Path(experiment_to_cpfe.__file__).resolve().parent.parent)}))'
    child=subprocess.run([sys.executable,'-B','-c',program,str(predicted/'predictions.npz'),str(Path(selected['path'])/'model.pt')],
        capture_output=True,text=True,timeout=180,env={**os.environ,'PYTHONPATH':str(backend_root)})
    save(context['output']/'model-readback.json',dict(exit_code=child.returncode,stdout=child.stdout,stderr=child.stderr))
    if child.returncode:
        raise RuntimeError('saved model failed independent process readback')
    report=json.loads((evaluated/'evaluation.json').read_text(encoding='utf-8'))
    from experiment_to_cpfe.mechanics.tensile import regression_metrics
    linear=next(c for c in candidates if c['id']=='linear')
    if selected['id']=='linear':
        linear_report=report
    else:
        linear_infer=dict(inference,checkpoint=str(Path(linear['path'])/'model.pt'),checkpoint_sha256=linear['model_sha256'])
        linear_output=invoke(context,'infer-baseline-linear','infer-surrogate',linear_infer,context['output']/'baselines/linear/inference')
        linear_eval=dict(evaluation,predictions=str(linear_output/'predictions.npz'))
        linear_evaluated=invoke(context,'evaluate-baseline-linear','evaluate-surrogate',linear_eval,context['output']/'baselines/linear/evaluation')
        linear_report=json.loads((linear_evaluated/'evaluation.json').read_text(encoding='utf-8'))
    baseline=regression_metrics(expected[:,0],np.full(len(expected),mean))
    comparison={'selected_model':selected['id'],'selection':'validation RMSE only',
        'target':'stress_reported','unit':'MPa','test_rows':len(expected),'test_specimens':1,
        'training_mean':baseline,'ols':linear_report['metrics']['stress_reported'],
        'selected':report['metrics']['stress_reported']}
    save(context['output']/'baseline-comparison.json',comparison)
    with (context['output']/'baseline-comparison.csv').open('x',encoding='utf-8-sig',newline='') as stream:
        writer=csv.writer(stream);writer.writerow(['method','target','unit','test_rows','test_specimens','rmse','mae','r2','bias'])
        for name,metric in [('training_mean',baseline),('OLS',comparison['ols']),('selected_'+selected['id'],comparison['selected'])]:
            writer.writerow([name,'stress_reported','MPa',len(expected),1,*[metric[k] for k in ['rmse','mae','r2','bias']]])
    return dict(training_executed=True,selected_model=selected['id'],model_candidates=candidates,
        test_metrics=report['metrics'],training_mean_MPa=mean,
        training_mean_test_baseline=baseline,ols_test_baseline=comparison['ols'],
        task_assessment=report['task_assessment'],new_process_readback=json.loads(child.stdout),
        test_row_identity_and_percent_conversion_exact=True)


def run_case(manifest, source_root, run_dir):
    started=time.perf_counter()
    context=prepare_run(manifest,source_root,run_dir)
    try:
        if manifest['case']=='dopamics':
            details=run_dopamics(context)
        else:
            entries=[normalize_record(context,r) for r in manifest['records']]
            details=dict(training_executed=False,total_rows=sum(e['rows'] for e in entries))
            if manifest['case']=='pcl':
                names=['time_min','temperature_C','stress_MPa'];units=['min','degC','MPa']
                output=grouped_build(context,entries,[dict(name=n,unit=u) for n,u in zip(names[:2],units[:2])],
                    [dict(name=names[2],unit=units[2])],{n:table_selector(n,u) for n,u in zip(names,units)})
                import numpy as np
                with np.load(output/'dataset.npz',allow_pickle=False) as packed:
                    for entry in entries:
                        record=next(r for r in manifest['records'] if r['id']==entry['id'])
                        values,_=read_numeric(context['paths'][entry['id']],record['reader'])
                        mask=packed['sample_ids']==entry['id']
                        if set(packed['splits'][mask])!={entry['role']}:
                            raise ValueError('dataset role mismatch')
                        indices=[record['reader']['fields'].index(n) for n in names]
                        np.testing.assert_array_equal(packed['features'][mask],np.asarray(values)[:,indices[:2]])
                        np.testing.assert_array_equal(packed['targets'][mask].reshape(-1),np.asarray(values)[:,indices[2]])
                details['task_assessment']=json.loads((output/'dataset.json').read_text(encoding='utf-8'))['task_assessment']
                details['numeric_dataset_readback_exact']=True
        after={str(p):digest(p) for p in context['sources'].values()}
        save(context['output']/'source-hashes-after.json',after)
        if after!=context['before']:
            raise ValueError('original source changed during run')
        result=dict(status='completed',case=manifest['case'],seconds=time.perf_counter()-started,
            source_hashes_unchanged=True,normalization_checks=context['readbacks'],
            solver_executed=False,third_party_code_executed=False,**details)
        limitations='\n'.join('- '+x for x in manifest.get('limitations',[]))
        text=f"# {manifest['case']} 外部案例运行结果\n\n工程流程已完成；原始文件摘要未变。科学有效性尚未建立。\n\n"
        text+=f"本次规范化记录数：{sum(e['rows'] for e in context['readbacks'])}（DOPAMICS 测试输入/真值分别核对，不作为独立观测数相加）。\n\n"
        if details['training_executed']:
            metrics=details['test_metrics']['stress_reported'];baseline=details['training_mean_test_baseline']
            ols=details['ols_test_baseline']
            text+=f"由验证集选择 {details['selected_model']}。测试仅一个原始试样，行数不等于独立试样数；未设置科学准确度验收门槛。\n\n"
            text+='|方法|RMSE (MPa)|MAE (MPa)|R²|bias (MPa)|\n|---|---:|---:|---:|---:|\n'
            for name,metric in [('训练均值',baseline),('OLS 线性',ols),('验证集所选 '+details['selected_model'],metrics)]:
                text+=f"|{name}|{metric['rmse']:.6g}|{metric['mae']:.6g}|{metric['r2']:.6g}|{metric['bias']:.6g}|\n"
            text+='\nRMSE/MAE 分别比较：较小的平方误差不保证较小的平均绝对误差。R² 为负表示平方误差超过以本测试真值均值为常数的参照；该测试均值仅用于定义指标，不是可部署的训练基线。\n\n'
            text+='预测与评价的可读表及中文摘要由标准入口生成，分别查看 inference/ 与 evaluation/；原始 JSON 收据继续保留。\n\n'
        else:
            text+='本例未训练或推理；FAIR 为完整原始通道接入，PCL 为包含预热的完整记录及数值建集。\n\n'
        text+='## 适用边界与下一步\n\n'+limitations+'\n\n先补充所列条件，再制定独立评价任务。不可把工程成功直接写成实验精度通过。\n'
        (context['output']/'REPORT.md').write_text(text,encoding='utf-8')
        (context['output']/'README.md').write_text('从 REPORT.md 开始；verification.json 为执行核验；commands/ 保留每条命令和耗时；configs/ 保留实际配置；source-hashes-before/after.json 核对原始来源。\n',encoding='utf-8')
        # The completed receipt is published only after every required delivery
        # artifact succeeds. A report write failure must remain a failed run.
        save(context['output']/'verification.json',result)
        return result
    except Exception as exc:
        save(context['output']/'failure.json',dict(status='failed',error=str(exc),seconds=time.perf_counter()-started))
        after={str(p):digest(p) for p in context['sources'].values()}
        if not (context['output']/'source-hashes-after.json').exists():
            save(context['output']/'source-hashes-after.json',after)
        raise


def main():
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, 'reconfigure'):
            stream.reconfigure(encoding='utf-8')
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--case',required=True,choices=['dopamics','fair-train','pcl'])
    parser.add_argument('--source-root',required=True,type=Path)
    parser.add_argument('--run-dir',required=True,type=Path)
    args=parser.parse_args()
    for name in ['OMP_NUM_THREADS','MKL_NUM_THREADS','OPENBLAS_NUM_THREADS','NUMEXPR_NUM_THREADS']:
        os.environ[name]='1'
    os.environ['CUDA_VISIBLE_DEVICES']='-1'
    os.environ['PYTHONDONTWRITEBYTECODE']='1'
    sys.path.insert(0,str(Path(__file__).resolve().parents[2]/'src'))
    manifest=json.loads((Path(__file__).parent/'manifests'/f'{args.case}.json').read_text(encoding='utf-8'))
    try:
        result=run_case(manifest,args.source_root,args.run_dir)
    except (ValueError,FileNotFoundError,FileExistsError,RuntimeError) as exc:
        print(str(exc),file=sys.stderr)
        return 1
    print(json.dumps(result,ensure_ascii=False,indent=2))
    return 0


if __name__=='__main__':
    raise SystemExit(main())
