"""Synthetic external-case tests: corrupt sources and mistaken rows must not pass."""
import hashlib
import importlib.util
import json
from pathlib import Path
import zipfile

import numpy as np
import pytest


def driver():
    path = Path(__file__).parents[2] / 'examples/external_cases/workflow.py'
    assert path.is_file(), 'portable external-case entry is missing'
    spec = importlib.util.spec_from_file_location('external_workflow', path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def fixture_manifest(tmp_path, *, archive=False):
    root = tmp_path / 'inputs'
    root.mkdir()
    data = b'StartOfData\n0;20;3\n1;21;4\n'
    checksum = hashlib.sha256(data).hexdigest()
    if archive:
        path = root / 'measurements.zip'
        with zipfile.ZipFile(path, 'w') as packed:
            packed.writestr('specimens/a.csv', data)
        source = {'path':path.name, 'sha256':hashlib.sha256(path.read_bytes()).hexdigest()}
        member = 'specimens/a.csv'
    else:
        path = root / 'a.csv'; path.write_bytes(data)
        source = {'path':path.name, 'sha256':checksum}
        member = None
    manifest = {'case':'pcl', 'doi':'synthetic', 'license':'synthetic fixture',
        'sources':[source], 'records':[{'id':'specimen-a','source':path.name,
            'member':member,'sha256':checksum,'output_name':'a.csv','role':'intake',
            'reader':{'format':'csv','encoding':'utf-8','delimiter':';','start_row':2,
                'fields':['time','temperature','stress'],'units':['min','degC','MPa'],
                'header_checks':[{'row':1,'column':0,'value':'StartOfData'}]}}]}
    return root, manifest


def test_changed_file_is_rejected_before_creating_output(tmp_path):
    module=driver(); root,manifest=fixture_manifest(tmp_path)
    (root/'a.csv').write_bytes(b'changed')
    out=tmp_path/'out'
    with pytest.raises(ValueError,match='SHA256'):
        module.prepare_run(manifest,root,out)
    assert not out.exists()


@pytest.mark.parametrize('mutation',['hash','identity'])
def test_original_identity_metadata_must_match_the_case_manifest(tmp_path,mutation):
    module=driver();root,manifest=fixture_manifest(tmp_path,archive=True)
    record=manifest['records'][0]
    metadata=b'Non-redundant ID (sample ID);Filename\nspecimen-a;Execution_essai1\n'
    with zipfile.ZipFile(root/'measurements.zip','a') as packed:
        packed.writestr('metadata.csv',metadata)
    manifest['sources'][0]['sha256']=hashlib.sha256((root/'measurements.zip').read_bytes()).hexdigest()
    manifest['identity_metadata']=dict(source='measurements.zip',member='metadata.csv',
                                     sha256=hashlib.sha256(metadata).hexdigest())
    record.update(trial=1,context={'Non-redundant ID (sample ID)':'specimen-a','Filename':'Execution_essai1'})
    if mutation=='hash':manifest['identity_metadata']['sha256']='0'*64
    else:record['context']['Filename']='Execution_essai2'
    with pytest.raises(ValueError,match='identity.*(SHA256|metadata|mapping)'):
        module.prepare_run(manifest,root,tmp_path/'out')
    assert not (tmp_path/'out').exists()


def test_existing_output_and_its_contents_are_preserved(tmp_path):
    module=driver(); root,manifest=fixture_manifest(tmp_path)
    out=tmp_path/'out'; out.mkdir(); marker=out/'keep.txt';marker.write_text('old')
    with pytest.raises(FileExistsError): module.prepare_run(manifest,root,out)
    assert marker.read_text()=='old' and list(out.iterdir())==[marker]


def test_archive_member_hash_checked_before_output(tmp_path):
    module=driver();root,manifest=fixture_manifest(tmp_path,archive=True)
    manifest['records'][0]['sha256']='0'*64
    with pytest.raises(ValueError,match='SHA256'):module.prepare_run(manifest,root,tmp_path/'out')
    assert not (tmp_path/'out').exists()


def test_manifest_cannot_read_outside_source_root(tmp_path):
    module=driver();root,manifest=fixture_manifest(tmp_path)
    manifest['sources'][0]['path']='../outside.csv'
    with pytest.raises(ValueError,match='relative'):module.prepare_run(manifest,root,tmp_path/'out')


def test_numeric_reader_preserves_negative_values_and_physical_rows(tmp_path):
    module=driver();root,manifest=fixture_manifest(tmp_path)
    p=root/'a.csv';p.write_text('StartOfData\n-1;20;3\n0;21;4\n')
    values,ids=module.read_numeric(p,manifest['records'][0]['reader'])
    np.testing.assert_array_equal(values,[[-1,20,3],[0,21,4]])
    assert ids==[2,3]


def test_instrument_header_padding_matches_existing_native_adapter(tmp_path):
    module=driver();root,manifest=fixture_manifest(tmp_path)
    p=root/'a.csv';p.write_text('StartOfData \n0;20;3\n')
    values,ids=module.read_numeric(p,manifest['records'][0]['reader'])
    assert values==[[0,20,3]] and ids==[2]


@pytest.mark.parametrize('body',['StartOfData\n','WrongHeader\n0;20;3\n',
    'StartOfData\n0;;3\n','StartOfData\n0;20;3;4\n','StartOfData\n0;nan;3\n'])
def test_empty_wrong_layout_and_nonfinite_values_rejected(tmp_path,body):
    module=driver();root,manifest=fixture_manifest(tmp_path)
    p=root/'a.csv';p.write_text(body)
    with pytest.raises(ValueError):module.read_numeric(p,manifest['records'][0]['reader'])


def test_xlsx_reader_rejects_formula_and_preserves_source_rows(tmp_path):
    module=driver()
    from openpyxl import Workbook
    p=tmp_path/'raw.xlsx'; book=Workbook();sheet=book.active;sheet.title='Sheet2'
    sheet.append(['title']);sheet.append(['Index','Load (N)']);sheet.append([0,-2]);sheet.append([1,3]);book.save(p)
    recipe={'format':'xlsx','sheet':'Sheet2','start_row':3,'fields':['index','force'],
            'units':['1','N'],'header_checks':[{'row':2,'column':1,'value':'Load (N)'}]}
    values,ids=module.read_numeric(p,recipe)
    np.testing.assert_array_equal(values,[[0,-2],[1,3]]);assert ids==[3,4]
    sheet.cell(3,2,'=1+1');book.save(p)
    with pytest.raises(ValueError):module.read_numeric(p,recipe)


def test_synthetic_normalization_runs_actual_cli_and_keeps_units(tmp_path):
    module=driver();root,manifest=fixture_manifest(tmp_path,archive=True)
    context=module.prepare_run(manifest,root,tmp_path/'out')
    record=manifest['records'][0]
    entry=module.normalize_record(context,record)
    assert entry['rows']==2 and entry['numeric_readback_exact']
    import h5py
    with h5py.File(Path(entry['hdf5']),'r') as handle:
        cols=handle['measured/measured_observations/columns']
        np.testing.assert_array_equal(cols['stress'][:],[3,4])
        np.testing.assert_array_equal(cols['source_row'][:],[2,3])
    receipt=json.loads((tmp_path/'out/commands/01-normalize-specimen-a.json').read_text())
    assert receipt['exit_code']==0


def test_dopamics_selection_uses_validation_only_and_linear_wins_tie():
    module=driver()
    candidates=[{'id':'linear','validation_rmse':2,'test_rmse':99},
                {'id':'mlp','validation_rmse':3,'test_rmse':0}]
    assert module.select_candidate(candidates)['id']=='linear'
    candidates[1]['validation_rmse']=2
    assert module.select_candidate(list(reversed(candidates)))['id']=='linear'


def test_selected_fields_keep_percentage_conversion_out_of_truth(tmp_path):
    module=driver();root,manifest=fixture_manifest(tmp_path)
    record=manifest['records'][0]
    record['reader']['fields']=['stress_reported','time','laser_strain_percent']
    record['reader']['units']=['MPa','s','%']
    context=module.prepare_run(manifest,root,tmp_path/'out')
    result=module.normalize_record(context,record,selected=['laser_strain_percent'],suffix='inputs-only')
    import h5py
    with h5py.File(result['hdf5'],'r') as handle:
        assert set(handle['measured/measured_observations/columns'])=={'laser_strain_percent','source_row'}
    selector=module.table_selector('laser_strain_percent','%',factor=.01)
    assert selector['conversion']['factor']==.01 and selector['source_unit']=='%'


def test_tiny_pcl_run_preserves_entire_records_and_group_roles(tmp_path):
    module=driver();root,manifest=fixture_manifest(tmp_path)
    manifest['sources']=[];manifest['records']=[]
    for n,role in enumerate(['train','validation','test'],1):
        path=root/f'specimen-{n}.csv'
        path.write_text(f'StartOfData\n0;{20+n};{n}\n1;{21+n};{n+1}\n')
        digest=hashlib.sha256(path.read_bytes()).hexdigest()
        manifest['sources'].append({'path':path.name,'sha256':digest})
        manifest['records'].append({'id':f'replicate-{n}','source':path.name,'sha256':digest,
            'output_name':path.name,'role':role,'reader':{'format':'csv','delimiter':';','encoding':'utf-8',
                'start_row':2,'fields':['time_min','temperature_C','stress_MPa'],
                'units':['min','degC','MPa'],'header_checks':[{'row':1,'column':0,'value':'StartOfData'}]}})
    result=module.run_case(manifest,root,tmp_path/'fresh')
    assert result['status']=='completed' and not result['training_executed']
    assert result['total_rows']==6 and result['source_hashes_unchanged']
    with np.load(tmp_path/'fresh/dataset/dataset.npz',allow_pickle=False) as package:
        assert package['targets'].reshape(-1).tolist()==[1,2,2,3,3,4]
        assert package['splits'].tolist()==['train','train','validation','validation','test','test']
        assert package['row_ids'].tolist()==['[2]','[3]']*3
    assert (tmp_path/'fresh/REPORT.md').is_file()
    assert (tmp_path/'fresh/source-hashes-after.json').is_file()


def test_report_write_failure_never_leaves_completed_verification(tmp_path,monkeypatch):
    module=driver();root,manifest=fixture_manifest(tmp_path)
    manifest['case']='fair-train'
    original=Path.write_text
    def fail_report(path,*args,**kwargs):
        if path.name=='REPORT.md':raise OSError('synthetic report disk failure')
        return original(path,*args,**kwargs)
    monkeypatch.setattr(Path,'write_text',fail_report)
    with pytest.raises(OSError,match='report disk failure'):
        module.run_case(manifest,root,tmp_path/'new')
    assert json.loads((tmp_path/'new/failure.json').read_text())['status']=='failed'
    assert not (tmp_path/'new/verification.json').exists()


def test_running_stage_reports_progress_to_stderr_without_polluting_stdout(tmp_path,capsys):
    module=driver();root,manifest=fixture_manifest(tmp_path)
    context=module.prepare_run(manifest,root,tmp_path/'new')
    module.normalize_record(context,manifest['records'][0])
    capture=capsys.readouterr()
    assert '开始' in capture.err and '完成' in capture.err
    assert capture.out==''


def test_feature_only_and_truth_only_preserve_same_sample_unit_system(tmp_path):
    from experiment_to_cpfe.datasets.hdf5 import read_hdf5

    module=driver();root,manifest=fixture_manifest(tmp_path)
    record=manifest['records'][0]
    context=module.prepare_run(manifest,root,tmp_path/'new')
    fields=record['reader']['fields']
    first=module.normalize_record(context,record,[fields[0]],'inputs-only')
    second=module.normalize_record(context,record,[fields[1]],'truth-only')
    first_sample=read_hdf5(first['hdf5']);second_sample=read_hdf5(second['hdf5'])
    expected=dict(zip(fields,record['reader']['units']))
    assert first_sample.metadata.unit_system == second_sample.metadata.unit_system == expected
    assert set(first_sample.tables['measured_observations'][0]) != set(second_sample.tables['measured_observations'][0])


def test_entry_uses_nonempty_cpu_guard_for_windows_children(tmp_path, monkeypatch, capsys):
    import os
    import sys

    module=driver()
    monkeypatch.setenv('CUDA_VISIBLE_DEVICES','original')
    monkeypatch.setattr(sys,'argv',['workflow.py','--case','pcl','--source-root',str(tmp_path),'--run-dir',str(tmp_path/'out')])
    observed={}
    def run(*args):
        observed['cuda']=os.environ.get('CUDA_VISIBLE_DEVICES')
        return {'status':'completed'}
    monkeypatch.setattr(module,'run_case',run)
    assert module.main()==0
    assert observed['cuda']=='-1'
